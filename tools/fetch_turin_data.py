"""Recreate public municipal geography, demographics and traffic snapshots.

Build-time only dependencies: pyshp, shapely, pyproj, openpyxl, xlrd.
Runtime planner uses only the checked-in processed JSON.
"""
from __future__ import annotations

import argparse
import hashlib
import io
import json
from pathlib import Path
from datetime import datetime, timezone
from urllib.request import Request, urlopen
import zipfile
import statistics
import unicodedata
import re
import xml.etree.ElementTree as ET
from html.parser import HTMLParser
from urllib.parse import urlencode

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "economics" / "data"
RAW = DATA / "raw"
SOURCES = {
    "boundaries": "https://www.istat.it/storage/cartografia/confini_amministrativi/generalizzati/2026/Limiti01012026_g.zip",
    "population": "https://www.regione.piemonte.it/web/media/54745/download",
    "commuting": "https://www.regione.piemonte.it/web/sites/default/files/media/documenti/2022-06/spostamenti_2018_2019.xls",
    "zones": "https://www.cittametropolitana.torino.it/istituzionale/che-cose-la-citta-metropolitana/le-zone-omogenee",
    "traffic": "https://opendata.5t.torino.it/get_fdt",
    "traffic_schema": "https://simone.5t.torino.it/ns/traffic_data.xsd",
}


def fetch(name: str, refresh=False):
    RAW.mkdir(parents=True, exist_ok=True)
    path = RAW / name
    if refresh or not path.exists():
        url = SOURCES[name]
        payload = None
        if name == "osm":
            url, query = url.split("?", 1)
            payload = query.encode("ascii")
        req = Request(url, data=payload, headers={"User-Agent": "OCCHIO-hackathon-public-data-research/1.0"})
        with urlopen(req, timeout=150) as response:
            content = response.read()
        path.write_bytes(content)
    return path.read_bytes()


class TextParser(HTMLParser):
    def __init__(self):
        super().__init__()
        self.parts = []

    def handle_data(self, data):
        self.parts.append(data.strip())


def normalize(name):
    value = unicodedata.normalize("NFKD", name).encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z0-9]", "", value)


ZONE_NAMES = {
    1: "Torino citta", 2: "Torino Ovest", 3: "Torino Sud", 4: "Torino Nord",
    5: "Pinerolese", 6: "Valli di Susa e Sangone", 7: "Ciriacese-Valli di Lanzo",
    8: "Canavese occidentale", 9: "Eporediese", 10: "Chivassese", 11: "Chierese-Carmagnolese",
}


def population_and_zones():
    import xlrd

    sheet = xlrd.open_workbook(file_contents=fetch("population")).sheet_by_index(0)
    municipalities = {}
    for i in range(2, sheet.nrows):
        row = sheet.row_values(i)
        if row[1] != "TO":
            continue
        municipalities[row[4]] = {
            "id": row[4], "name": row[5], "population": int(row[6]), "population_year": 2024,
            "population_date": "2024-12-31", "area_km2": float(row[8]),
            "area_year": 2021, "density_per_km2": round(float(row[7]), 2),
            "urbanization_class": int(row[11]),
        }
    assert len(municipalities) == 312
    names = {normalize(m["name"]): m for m in municipalities.values()}
    parser = TextParser()
    parser.feed(fetch("zones").decode("utf-8"))
    text = " ".join(parser.parts)
    aliases = {
        "Venaria": "Venaria Reale", "Castagnole P.te": "Castagnole Piemonte",
        "San Germano C.": "San Germano Chisone", "San Secondo di P.": "San Secondo di Pinerolo",
        "Cesana T.se": "Cesana Torinese", "Coassolo T.se": "Coassolo Torinese", "Borgone di Susa": "Borgone Susa",
        "San Francesco al C.": "San Francesco al Campo", "San Maurizio C.se": "San Maurizio Canavese",
        "Baldissero C.se": "Baldissero Canavese", "Colleretto C.": "Colleretto Castelnuovo",
        "Forno C.se": "Forno Canavese", "San Colombano B.": "San Colombano Belmonte",
        "San Giorgio C.se": "San Giorgio Canavese", "San Giusto C.se": "San Giusto Canavese",
        "Barone C.se": "Barone Canavese", "Borgofranco": "Borgofranco d'Ivrea",
        "Candi,a C.se": "Candia Canavese", "Cossano C.se": "Cossano Canavese",
        "Fiorano C.se": "Fiorano Canavese", "San Martino C.se": "San Martino Canavese",
        "Montaldo T.se": "Montaldo Torinese", "Moriondo T.se": "Moriondo Torinese",
        "Pecetto T.se": "Pecetto Torinese",
    }
    names[normalize("Torino")]["zone"] = 1
    for zone in range(2, 12):
        if zone == 4:
            continue
        start = text.index(f"Zona {zone} ")
        end = text.find(f"Zona {zone+1} ", start) if zone < 11 else text.find("Ultimo aggiornamento", start)
        portion = text[start:end]
        listed = re.split(r"Comuni\s*:", portion, maxsplit=1)[1].strip().rstrip(".")
        # The source contains a typographical comma inside Candia's name.
        listed = listed.replace("Candi,a C.se", "Candia Canavese")
        for name in listed.split(","):
            name = name.strip().rstrip(".")
            match = names.get(normalize(name)) or names.get(normalize(aliases.get(name, aliases.get(name+".", ""))))
            if not match:
                raise ValueError(f"Unmatched municipality in official zone {zone}: {name}")
            match["zone"] = zone
    remaining = [m for m in municipalities.values() if "zone" not in m]
    assert len(remaining) == 8, [m["name"] for m in remaining]
    for m in remaining:
        m["zone"] = 4
    for m in municipalities.values():
        m["zone_name"] = ZONE_NAMES[m["zone"]]
    expected = {1: 1, 2: 14, 3: 18, 4: 8, 5: 45, 6: 40, 7: 40, 8: 46, 9: 54, 10: 24, 11: 22}
    assert {z: sum(m["zone"] == z for m in municipalities.values()) for z in expected} == expected
    commuting = xlrd.open_workbook(file_contents=fetch("commuting")).sheet_by_index(0)
    for i in range(5, commuting.nrows):
        row = commuting.row_values(i)
        if row[0] not in municipalities:
            continue
        m = municipalities[row[0]]
        m.update(commuters_work=int(row[17]), commuters_study=int(row[18]),
                 commuters_total=int(row[19]), commuters_internal=int(row[13]),
                 commuters_outbound=int(row[16]), commuting_year=2019)
    assert all(m["commuters_work"] + m["commuters_study"] == m["commuters_total"]
               and m["commuters_internal"] + m["commuters_outbound"] == m["commuters_total"]
               for m in municipalities.values())
    return municipalities


def geometries():
    import shapefile
    from pyproj import CRS, Transformer
    from shapely.geometry import shape, mapping
    from shapely.ops import transform

    polygons = {}
    with zipfile.ZipFile(io.BytesIO(fetch("boundaries"))) as archive:
        shp = next(x for x in archive.namelist() if "/Com" in x and x.endswith(".shp"))
        reader = shapefile.Reader(shp=io.BytesIO(archive.read(shp)), dbf=io.BytesIO(archive.read(shp[:-4]+".dbf")), encoding="utf-8")
        crs = CRS.from_wkt(archive.read(shp[:-4]+".prj").decode())
        convert = Transformer.from_crs(crs, 4326, always_xy=True).transform
        for record in reader.iterShapeRecords():
            props = record.record.as_dict()
            if int(props["COD_CM"]) != 201:
                continue
            geom = shape(record.shape.__geo_interface__)
            if not geom.is_valid:
                geom = geom.buffer(0)
            # Retain detailed official geometry for point assignment; simplify map only.
            polygons[props["PRO_COM_T"]] = transform(convert, geom)
    assert len(polygons) == 312
    return polygons


_SPATIAL_INDEX = None


def assign_point(polygons, longitude, latitude):
    from shapely.geometry import Point
    from shapely import STRtree

    global _SPATIAL_INDEX
    if _SPATIAL_INDEX is None or _SPATIAL_INDEX[0] is not polygons:
        _SPATIAL_INDEX = (polygons, list(polygons), STRtree(list(polygons.values())))
    point = Point(longitude, latitude)
    matches = _SPATIAL_INDEX[2].query(point, predicate="intersects")
    if len(matches):
        return _SPATIAL_INDEX[1][int(matches[0])]
    return None


def traffic_snapshot(municipalities, polygons, refresh=False):
    root = ET.fromstring(fetch("traffic", refresh))
    points = []
    grouped = {}
    for row in root:
        if row.tag.split("}")[-1] != "FDT_data":
            continue
        a = row.attrib
        longitude, latitude = float(a["lng"]), float(a["lat"])
        code = assign_point(polygons, longitude, latitude)
        if not code:
            continue
        flow = next(iter(row)).attrib
        road = a.get("Road_name", "")
        item = {"municipality_id": code, "longitude": longitude, "latitude": latitude,
                "road": road, "direction": a.get("direction"), "offset_m": float(a.get("offset", 0)),
                "location_code": a.get("lcd1"), "flow_vph": float(flow["flow"]),
                "speed_kph": float(flow["speed"]), "accuracy": float(a.get("accuracy", 0)),
                "period_minutes": float(a.get("period", 5))}
        # Motorway and airport-link sections cannot calibrate pedestrian crossing queues.
        item["used_local_benchmark"] = item["accuracy"] > 0 and not re.match(r"^(A\d|R\d|Tangenziale)", road, re.I)
        points.append(item)
        if item["used_local_benchmark"]:
            grouped.setdefault(code, []).append(item["flow_vph"])
    for code, m in municipalities.items():
        local = grouped.get(code, [])
        m["traffic_vph"] = round(statistics.median(local), 1) if local else None
        m["traffic_counter_count"] = len(local)
        m["traffic_source"] = "5T single 5-minute directional road-section median" if local else None
    output = {"source": SOURCES["traffic"], "generation_time": root.attrib["generation_time"],
              "start_time": root.attrib["start_time"], "end_time": root.attrib["end_time"],
              "flow_unit": "vehicles/hour", "records": points,
              "caveat": "One Saturday afternoon snapshot, not annual average traffic. Directional section medians are not crossing-specific counts. Zero-accuracy counters and motorway/raccordo sections excluded from local benchmarks."}
    (DATA / "traffic_snapshot.json").write_text(json.dumps(output, ensure_ascii=True, indent=2), encoding="utf-8")


def osm_inventory(municipalities, polygons, refresh=False):
    bbox = (44.70, 6.55, 45.70, 8.30)
    query = f'[out:json][timeout:180];node["highway"="crossing"]({",".join(str(x) for x in bbox)});out body;'
    url = "https://overpass-api.de/api/interpreter?" + urlencode({"data": query})
    SOURCES["osm"] = url
    content = json.loads(fetch("osm", refresh))
    for m in municipalities.values():
        m["mapped_crossings"] = 0
        m["signalized_crossings"] = 0
    points = []
    for row in content["elements"]:
        code = assign_point(polygons, row["lon"], row["lat"])
        if code is None:
            continue
        tags = row.get("tags", {})
        signals = tags.get("crossing") == "traffic_signals" or tags.get("crossing:signals") == "yes"
        municipalities[code]["mapped_crossings"] += 1
        municipalities[code]["signalized_crossings"] += int(signals)
        points.append({"osm_id": row["id"], "municipality_id": code, "longitude": row["lon"],
                       "latitude": row["lat"], "signalized": signals})
    (DATA / "crossings.json").write_text(json.dumps({"timestamp": content["osm3s"]["timestamp_osm_base"],
             "license": "ODbL 1.0", "attribution": "OpenStreetMap contributors", "records": points,
             "caveat": "Mapped highway=crossing nodes, not a complete asset register or certified installable sites. Separate lane nodes may belong to one crossing complex. No deduplication into installation sites."}, ensure_ascii=True, separators=(",", ":")), encoding="utf-8")


def build(include_osm=False, refresh_traffic=False):
    from shapely.geometry import mapping

    DATA.mkdir(parents=True, exist_ok=True)
    municipalities = population_and_zones()
    polygons = geometries()
    assert set(municipalities) == set(polygons)
    traffic_snapshot(municipalities, polygons, refresh_traffic)
    for m in municipalities.values():
        m.update(mapped_crossings=None, signalized_crossings=None)
    if include_osm or (RAW / "osm").exists():
        osm_inventory(municipalities, polygons)
    assert all(m["population"] > 0 and m["area_km2"] > 0 for m in municipalities.values())
    assert all(m["mapped_crossings"] is None or 0 <= m["signalized_crossings"] <= m["mapped_crossings"]
               for m in municipalities.values())
    features = []
    for code in sorted(municipalities):
        geometry = polygons[code].simplify(0.0004, preserve_topology=True)
        m = municipalities[code]
        center = polygons[code].representative_point()
        m["center"] = [round(center.x, 6), round(center.y, 6)]
        features.append({"type": "Feature", "id": code, "properties": m, "geometry": mapping(geometry)})
    output = {"type": "FeatureCollection", "name": "Metropolitan City of Turin",
              "features": features}
    (DATA / "municipalities.geojson").write_text(json.dumps(output, ensure_ascii=True, separators=(",", ":")), encoding="utf-8")
    manifest = {"dataset_version": "turin-2026-boundaries-2024-population-v1",
                "retrieved_at": datetime.now(timezone.utc).isoformat(), "municipalities": 312,
                "population": sum(x["population"] for x in municipalities.values()),
                "sources": [], "caveats": [
                    "Official boundaries2026 and population31December2024 are different reference years.",
                    "Commuter counts2019 describe resident people, not total trips or pedestrian crossing demand.",
                    "Zona4 assigned as remainder of the official11-zone partition; source page omits Leini from its spokesperson listing.",
                    "Map geometry is additionally simplified0.0004degrees; statistical boundaries are not cadastral/legal boundaries.",
                    "Observed crossing nodes are not validated deployment sites. Missing OSM data is null, not zero.",
                    "5T counters do not cover all municipalities, and one snapshot cannot calibrate annual benefits."]}
    descriptions = {
        "boundaries": ("ISTAT administrative boundaries", "2026-01-01", "CC BY 4.0; attribute ISTAT", "Official generalized UTM32N shapefile, transformed to WGS84 longitude/latitude."),
        "population": ("Regione Piemonte / ISTAT resident population", "2024-12-31", "Public statistical data; attribute Regione Piemonte and ISTAT", "Definitive continuous-census population. First worksheet only; sorted worksheet contains inconsistent numeric municipality codes and is not used."),
        "commuting": ("Regione Piemonte / ISTAT commuting", "2019", "Public statistical data; attribute Regione Piemonte and ISTAT", "Resident daily work/study commuters, all transport modes."),
        "zones": ("CMTO homogeneous zones", "2025-11-17 page update", "Official administrative grouping; attribute CMTO", "Source spelling abbreviations normalized to official municipality names."),
        "traffic": ("5T measured traffic", "Single5-minute snapshot", "IODL 2.0 per aperTO municipal data catalog; attribute Citta di Torino /5T", "Indicative flow, no guarantee of precise vehicle counts."),
        "osm": ("OpenStreetMap crossing nodes via Overpass", "Snapshot", "ODbL 1.0", "Node inventory incomplete, road-crossing complexes can have multiple nodes."),
    }
    for name, url in SOURCES.items():
        path = RAW / name
        if not path.exists() or name not in descriptions:
            continue
        title, reference, license_, caveat = descriptions[name]
        manifest["sources"].append({"id": name, "title": title, "url": url, "reference": reference,
                                     "license": license_, "caveat": caveat,
                                     "retrieved_at": datetime.fromtimestamp(path.stat().st_mtime, timezone.utc).isoformat(),
                                     "sha256": hashlib.sha256(path.read_bytes()).hexdigest()})
    manifest["sources"].append({
        "id": "pums_calibration", "title": "CMTO PUMS model calibration report",
        "url": "https://www.cittametropolitana.torino.it/sites/default/files/pagina/allegati/trasporti/PUMS/ALLEGATI/ALLEGATO_I_-_Rapporto_di_calibrazione_del_modello_v12.pdf",
        "reference": "2019 pre-COVID baseline; approved July2022",
        "license": "Official public planning report; attribute CMTO /META /TRASPOL",
        "caveat": "Context only. Its validated regional demand model and planned 2030 intervention effects are not the OCCHIO model or causal OCCHIO effectiveness evidence.",
    })
    (DATA / "manifest.json").write_text(json.dumps(manifest, indent=2, ensure_ascii=True), encoding="utf-8")
    print(json.dumps({"municipalities": len(features), "population": manifest["population"], "map_bytes": (DATA/"municipalities.geojson").stat().st_size, "zone_counts": {z:sum(x["zone"]==z for x in municipalities.values()) for z in ZONE_NAMES}}))


def inspect():
    import shapefile
    import openpyxl
    import xlrd

    content = fetch("boundaries")
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        shp = next(x for x in archive.namelist() if "/Com" in x and x.endswith(".shp"))
        reader = shapefile.Reader(shp=io.BytesIO(archive.read(shp)), dbf=io.BytesIO(archive.read(shp[:-4]+".dbf")))
        print("SHAPE", shp, reader.fields, reader.record(0))
    content = fetch("population")
    if content.startswith(b"PK"):
        workbook = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        for sheet in workbook:
            print("POP", sheet.title, sheet.max_row, sheet.max_column)
            for row in list(sheet.iter_rows(values_only=True))[:12]:
                print(row[:20])
    else:
        workbook = xlrd.open_workbook(file_contents=content)
        for sheet in workbook.sheets():
            print("POP", sheet.name, sheet.nrows, sheet.ncols)
            for row in range(min(12, sheet.nrows)):
                print(sheet.row_values(row)[:20])
    content = fetch("traffic")
    print("TRAFFIC", content[:6500].decode("utf-8", errors="replace"))


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("--inspect", action="store_true")
    parser.add_argument("--osm", action="store_true")
    parser.add_argument("--refresh-traffic", action="store_true")
    args = parser.parse_args()
    if args.inspect:
        inspect()
    else:
        build(args.osm, args.refresh_traffic)
